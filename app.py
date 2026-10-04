import os
import io
import json
import time
import random
import string
import logging
import base64
import mimetypes
import re
from functools import wraps

import requests
from flask import Flask, request, jsonify, render_template_string, session, redirect, url_for
from flask_cors import CORS
import telebot
from telebot import types

# Optional document readers
try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    from captcha.image import ImageCaptcha
except Exception:
    ImageCaptcha = None

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ============================================================
# Configuration
# ============================================================
TOKEN = (os.environ.get("TELEGRAM_BOT_TOKEN") or os.environ.get("BOT_TOKEN", "")).strip()

raw_admin = os.environ.get("ADMIN_ID", "0").strip()
ADMIN = int(raw_admin) if raw_admin.isdigit() else 0

CLOUDFLARE_ACCOUNT_ID = os.environ.get("CLOUDFLARE_ACCOUNT_ID", "").strip()
CLOUDFLARE_API_TOKEN = os.environ.get("CLOUDFLARE_API_TOKEN", "").strip()
CLOUDFLARE_MODEL = os.environ.get("CLOUDFLARE_MODEL", "@cf/meta/llama-3-8b-instruct").strip()
CLOUDFLARE_VISION_MODEL = os.environ.get(
    "CLOUDFLARE_VISION_MODEL", "@cf/meta/llama-3.2-11b-vision-instruct"
).strip()

DASHBOARD_PASSWORD = os.environ.get("DASHBOARD_PASSWORD", "admin123").strip()
SECRET_KEY = os.environ.get("SECRET_KEY", "").strip()

APP_NAME = os.environ.get("FLY_APP_NAME", "")
DEFAULT_HOST = f"https://{APP_NAME}.fly.dev" if APP_NAME else "http://localhost:8080"
WEBHOOK_HOST = os.environ.get("WEBHOOK_HOST", DEFAULT_HOST).rstrip("/")
WEBHOOK_PATH = f"/webhook/{TOKEN}" if TOKEN else "/webhook"
WEBHOOK_URL = f"{WEBHOOK_HOST}{WEBHOOK_PATH}"

BOT_NAME = os.environ.get("BOT_NAME", "مريم محمد جاسم الياسري💙🎀")
ENDPOINT_PATH = "/api/bot"

DB_DIR = os.environ.get("DB_DIR", os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(DB_DIR, "bot_data.json")

MAX_DOWNLOAD_MB = int(os.environ.get("MAX_DOWNLOAD_MB", "15"))
MAX_TEXT_CHARS = int(os.environ.get("MAX_TEXT_CHARS", "60000"))

os.makedirs(DB_DIR, exist_ok=True)

app = Flask(__name__)
app.secret_key = SECRET_KEY or os.urandom(32)
CORS(app)

bot = telebot.TeleBot(TOKEN, threaded=False) if TOKEN else None

DEFAULT_DB = {
    "users": [],
    "notify": True,
    "superadmins": [ADMIN] if ADMIN else [],
    "api_keys": {},
    "pending_admins": [],
    "approved_admins": [ADMIN] if ADMIN else [],
    "verify_enabled": False,
}

conversations = {}
pending_verifications = {}

# ============================================================
# Database Functions
# ============================================================
def load_db():
    try:
        if not os.path.exists(DB):
            save_db(DEFAULT_DB.copy())
            return DEFAULT_DB.copy()
        with open(DB, "r", encoding="utf-8") as f:
            data = json.load(f)
        for key, value in DEFAULT_DB.items():
            if key not in data:
                data[key] = value.copy() if isinstance(value, (list, dict)) else value
        if ADMIN and ADMIN not in data.get("superadmins", []):
            data.setdefault("superadmins", []).append(ADMIN)
        if ADMIN and ADMIN not in data.get("approved_admins", []):
            data.setdefault("approved_admins", []).append(ADMIN)
        return data
    except Exception as exc:
        logger.exception("DB LOAD ERROR: %s", exc)
        return DEFAULT_DB.copy()


def save_db(data):
    try:
        tmp = DB + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        os.replace(tmp, DB)
    except Exception as exc:
        logger.exception("DB SAVE ERROR: %s", exc)


def is_superadmin(uid):
    if not uid:
        return False
    if ADMIN and int(uid) == ADMIN:
        return True
    data = load_db()
    return int(uid) in [int(x) for x in data.get("superadmins", [])]


def is_approved_admin(uid):
    if not uid:
        return False
    if is_superadmin(uid):
        return True
    data = load_db()
    return int(uid) in [int(x) for x in data.get("approved_admins", [])]


def generate_api_key():
    chars = string.ascii_letters + string.digits
    return "FM.AI_" + "".join(random.choice(chars) for _ in range(30))


def get_dashboard_api_key():
    data = load_db()
    keys = list(data.get("api_keys", {}).keys())
    if keys:
        return keys[0]
    key = generate_api_key()
    data.setdefault("api_keys", {})[key] = ADMIN or 0
    save_db(data)
    return key

# ============================================================
# AI Functions
# ============================================================
def cloudflare_url(model):
    if not CLOUDFLARE_ACCOUNT_ID or not model:
        return None
    return f"https://api.cloudflare.com/client/v4/accounts/{CLOUDFLARE_ACCOUNT_ID}/ai/run/{model}"


def ask_ai(text, system_prompt=None, conversation_id=None, model=None, extra_messages=None):
    api_url = cloudflare_url(model or CLOUDFLARE_MODEL)
    if not api_url or not CLOUDFLARE_API_TOKEN:
        return "⚠️ الذكاء الاصطناعي غير مضبوط حالياً."

    system_prompt = system_prompt or (
        f"اسمك {BOT_NAME}. أجب بالعربية بدقة ووضوح."
    )
    headers = {
        "Authorization": f"Bearer {CLOUDFLARE_API_TOKEN}",
        "Content-Type": "application/json",
    }
    messages = [{"role": "system", "content": system_prompt}]
    if conversation_id and conversation_id in conversations:
        messages.extend(conversations[conversation_id][-20:])
    if extra_messages:
        messages.extend(extra_messages)
    messages.append({"role": "user", "content": text})
    payload = {"messages": messages, "max_tokens": 1800}

    try:
        response = requests.post(api_url, headers=headers, json=payload, timeout=90)
        data = response.json()
        if not data.get("success"):
            logger.error("Cloudflare error: %s", data.get("errors"))
            return None
        result = data.get("result") or {}
        reply = result.get("response")
        if not reply and result.get("choices"):
            reply = result["choices"][0].get("message", {}).get("content")
        if reply and conversation_id:
            conversations.setdefault(conversation_id, [])
            conversations[conversation_id].extend([
                {"role": "user", "content": text},
                {"role": "assistant", "content": reply},
            ])
            conversations[conversation_id] = conversations[conversation_id][-20:]
        return reply
    except Exception as exc:
        logger.exception("Cloudflare AI ERROR: %s", exc)
        return None


def ask_vision(prompt, image_bytes, mime_type="image/jpeg"):
    api_url = cloudflare_url(CLOUDFLARE_VISION_MODEL)
    if not api_url or not CLOUDFLARE_API_TOKEN:
        return None
    b64 = base64.b64encode(image_bytes).decode("ascii")
    image_url = f"data:{mime_type};base64,{b64}"
    headers = {
        "Authorization": f"Bearer {CLOUDFLARE_API_TOKEN}",
        "Content-Type": "application/json",
    }
    payload = {
        "messages": [
            {"role": "system", "content": f"أنت مساعد اسمه {BOT_NAME}. حلل الصور بالعربية بدقة."},
            {"role": "user", "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": image_url}},
            ]},
        ],
        "max_tokens": 1800,
    }
    try:
        response = requests.post(api_url, headers=headers, json=payload, timeout=90)
        data = response.json()
        if not data.get("success"):
            return None
        result = data.get("result") or {}
        return result.get("response") or (result.get("choices", [{}])[0].get("message", {}).get("content"))
    except Exception as exc:
        logger.exception("Vision ERROR: %s", exc)
        return None

# ============================================================
# File Helpers
# ============================================================
CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css", ".scss",
    ".json", ".xml", ".yaml", ".yml", ".md", ".txt", ".sql", ".php", ".java",
    ".kt", ".c", ".h", ".cpp", ".cs", ".go", ".rs", ".rb", ".dart", ".sh",
}

def safe_filename(name):
    name = os.path.basename(name or "file")
    return re.sub(r"[^\w.\- ()\[\]]", "_", name, flags=re.UNICODE)


def extract_text_from_file(filename, raw):
    ext = os.path.splitext(filename.lower())[1]
    if ext in CODE_EXTENSIONS or (not ext and len(raw) < 2_000_000):
        for encoding in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
            try:
                return raw.decode(encoding), "text/code"
            except Exception:
                pass
    if ext == ".pdf" and PdfReader:
        try:
            reader = PdfReader(io.BytesIO(raw))
            return "\n\n".join([(p.extract_text() or "") for p in reader.pages[:30]]), "PDF"
        except Exception as e:
            return f"تعذر قراءة PDF: {e}", "PDF"
    if ext == ".docx" and DocxDocument:
        try:
            doc = DocxDocument(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs), "DOCX"
        except Exception as e:
            return f"تعذر قراءة DOCX: {e}", "DOCX"
    if ext == ".xlsx" and openpyxl:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=100, values_only=True):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(parts), "XLSX"
        except Exception as e:
            return f"تعذر قراءة XLSX: {e}", "XLSX"
    return None, "unsupported"


def download_telegram_file(file_id):
    if not bot:
        raise RuntimeError("Bot not configured")
    info = bot.get_file(file_id)
    size = getattr(info, "file_size", 0) or 0
    if size > MAX_DOWNLOAD_MB * 1024 * 1024:
        raise ValueError(f"حجم الملف أكبر من الحد المسموح ({MAX_DOWNLOAD_MB}MB).")
    return bot.download_file(info.file_path), info.file_path

# ============================================================
# Keyboards
# ============================================================
def main_menu(is_super=False):
    kb = types.InlineKeyboardMarkup(row_width=2)
    kb.add(
        types.InlineKeyboardButton("💬 محادثة", callback_data="chat"),
        types.InlineKeyboardButton("❓ مساعدة", callback_data="help")
    )
    kb.add(types.InlineKeyboardButton("👨‍💻 المطورين", callback_data="dev"))
    if is_super:
        kb.add(types.InlineKeyboardButton("👑 المشرفين", callback_data="superadmin"))
    return kb


def back_button():
    kb = types.InlineKeyboardMarkup()
    kb.add(types.InlineKeyboardButton("🔙 رجوع", callback_data="back"))
    return kb


def dev_menu():
    kb = types.InlineKeyboardMarkup(row_width=1)
    kb.add(
        types.InlineKeyboardButton("🔑 إنشاء رمز API", callback_data="gen_api"),
        types.InlineKeyboardButton("📖 دليل المطورين", callback_data="dev_guide"),
        types.InlineKeyboardButton("🔙 رجوع", callback_data="back")
    )
    return kb


def superadmin_menu():
    kb = types.InlineKeyboardMarkup(row_width=1)
    kb.add(
        types.InlineKeyboardButton("📢 إذاعة", callback_data="broadcast_menu"),
        types.InlineKeyboardButton("👥 المستخدمين", callback_data="users_count"),
        types.InlineKeyboardButton("🔙 رجوع", callback_data="back")
    )
    return kb


def broadcast_type_menu():
    kb = types.InlineKeyboardMarkup(row_width=2)
    for label, key in [("📝 نص", "text"), ("🖼️ صورة", "photo"), ("🎥 فيديو", "video")]:
        kb.add(types.InlineKeyboardButton(label, callback_data=f"bcast_{key}"))
    kb.add(types.InlineKeyboardButton("🔙 رجوع", callback_data="superadmin"))
    return kb


def approval_buttons(uid):
    kb = types.InlineKeyboardMarkup(row_width=2)
    kb.add(
        types.InlineKeyboardButton("✅ موافقة", callback_data=f"approve_{uid}"),
        types.InlineKeyboardButton("❌ رفض", callback_data=f"reject_{uid}")
    )
    return kb

# ============================================================
# Flask Endpoints
# ============================================================
@app.route(WEBHOOK_PATH, methods=["POST"])
def webhook():
    if not bot:
        return "Bot disabled", 503
    try:
        raw = request.get_data().decode("utf-8")
        update = types.Update.de_json(raw)
        bot.process_new_updates([update])
        return "OK", 200
    except Exception as exc:
        logger.exception("Webhook processing error: %s", exc)
        return "Error", 500


@app.route("/")
def index():
    return jsonify({
        "status": "online",
        "bot": BOT_NAME,
        "webhook": WEBHOOK_URL,
        "admin_configured": bool(ADMIN)
    })


@app.route("/setwebhook")
def set_webhook():
    if not bot:
        return jsonify({"ok": False, "error": "Token missing"}), 503
    try:
        bot.remove_webhook()
        time.sleep(0.5)
        res = bot.set_webhook(url=WEBHOOK_URL)
        return jsonify({"ok": True, "webhook_url": WEBHOOK_URL, "result": bool(res)})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

# ============================================================
# Telegram Handlers
# ============================================================
def remember_user(user):
    data = load_db()
    if user.id not in data["users"]:
        data["users"].append(user.id)
        save_db(data)
        if data.get("notify", True):
            for aid in data.get("superadmins", []):
                try:
                    bot.send_message(aid, f"🆕 مستخدم جديد: {user.first_name}\n🆔: `{user.id}`", parse_mode="Markdown")
                except Exception:
                    pass


def send_long(chat_id, text):
    text = text or "لا يوجد رد."
    for i in range(0, len(text), 4000):
        bot.send_message(chat_id, text[i:i + 4000])


if bot:
    @bot.message_handler(commands=["start"])
    def cmd_start(message):
        remember_user(message.from_user)
        user_is_super = is_superadmin(message.from_user.id)
        welcome = (
            f"🌟 أهلاً بك {message.from_user.first_name} في {BOT_NAME}!\n\n"
            "🤖 أستطيع مساعدتك في الإجابة عن الأسئلة، وتحليل الكود والملفات والصور."
        )
        bot.send_message(
            message.chat.id,
            welcome,
            reply_markup=main_menu(is_super=user_is_super)
        )

    # معالج الأزرار الشفافة المصحح
    @bot.callback_query_handler(func=lambda call: True)
    def handle_all_callbacks(call):
        try:
            bot.answer_callback_query(call.id)
        except Exception:
            pass

        if not call.message:
            return

        chat_id = call.message.chat.id
        msg_id = call.message.message_id
        data = call.data
        user_id = call.from_user.id
        user_is_super = is_superadmin(user_id)
        db_data = load_db()

        try:
            if data == "chat":
                bot.edit_message_text(
                    "💬 أرسل استفسارك الآن وسأجيبك فوراً.",
                    chat_id=chat_id,
                    message_id=msg_id,
                    reply_markup=back_button()
                )
            elif data == "help":
                bot.edit_message_text(
                    "📖 يمكنك إرسال نصوص، صور، أو ملفات برمجية لتحليلها.",
                    chat_id=chat_id,
                    message_id=msg_id,
                    reply_markup=back_button()
                )
            elif data == "dev":
                if not is_approved_admin(user_id):
                    if user_id not in db_data.get("pending_admins", []):
                        db_data.setdefault("pending_admins", []).append(user_id)
                        save_db(db_data)
                        for aid in db_data.get("superadmins", []):
                            try:
                                bot.send_message(
                                    aid,
                                    f"📩 طلب انضمام مطور:\n👤: {call.from_user.first_name}\n🆔: `{user_id}`",
                                    parse_mode="Markdown",
                                    reply_markup=approval_buttons(user_id)
                                )
                            except Exception:
                                pass
                    bot.edit_message_text(
                        "📩 تم إرسال طلبك للإدارة للموافقة.",
                        chat_id=chat_id,
                        message_id=msg_id,
                        reply_markup=back_button()
                    )
                else:
                    bot.edit_message_text(
                        "👨‍💻 أهلاً بك في لوحة المطورين:",
                        chat_id=chat_id,
                        message_id=msg_id,
                        reply_markup=dev_menu()
                    )
            elif data == "gen_api":
                if not is_approved_admin(user_id):
                    bot.send_message(chat_id, "⛔ ليس لديك صلاحية مطور بعد.")
                    return
                key = generate_api_key()
                db_data.setdefault("api_keys", {})[key] = user_id
                save_db(db_data)
                bot.edit_message_text(
                    f"🔑 رمز الـ API الجديد الخاص بك:\n\n`{key}`",
                    chat_id=chat_id,
                    message_id=msg_id,
                    parse_mode="Markdown",
                    reply_markup=dev_menu()
                )
            elif data == "dev_guide":
                bot.edit_message_text(
                    f"📖 رابط واجهة الـ API:\n\nPOST {WEBHOOK_HOST}{ENDPOINT_PATH}\nAuthorization: Bearer YOUR_KEY\nJSON: {{\"message\":\"hi\"}}",
                    chat_id=chat_id,
                    message_id=msg_id,
                    reply_markup=dev_menu()
                )
            elif data == "superadmin":
                if not user_is_super:
                    bot.send_message(chat_id, "⛔ هذا القسم مخصص للمشرفين فقط.")
                    return
                bot.edit_message_text(
                    "👑 لوحة تحكم المشرفين:",
                    chat_id=chat_id,
                    message_id=msg_id,
                    reply_markup=superadmin_menu()
                )
            elif data == "users_count":
                if user_is_super:
                    bot.edit_message_text(
                        f"👥 عدد المستخدمين: {len(db_data.get('users', []))}",
                        chat_id=chat_id,
                        message_id=msg_id,
                        reply_markup=superadmin_menu()
                    )
            elif data == "broadcast_menu":
                if user_is_super:
                    bot.edit_message_text(
                        "📢 حدد نوع الإذاعة:",
                        chat_id=chat_id,
                        message_id=msg_id,
                        reply_markup=broadcast_type_menu()
                    )
            elif data.startswith("bcast_"):
                if not user_is_super:
                    return
                btype = data[6:]
                sent_msg = bot.edit_message_text(
                    f"📝 أرسل الآن رسالة الإذاعة ({btype}):",
                    chat_id=chat_id,
                    message_id=msg_id
                )
                bot.register_next_step_handler(sent_msg, lambda m: do_broadcast(m, btype))
            elif data.startswith("approve_"):
                uid = int(data[8:])
                if not user_is_super:
                    return
                db_data.setdefault("approved_admins", []).append(uid)
                if uid in db_data.get("pending_admins", []):
                    db_data["pending_admins"].remove(uid)
                save_db(db_data)
                bot.edit_message_text("✅ تمت الموافقة بنجاح.", chat_id=chat_id, message_id=msg_id)
                try:
                    bot.send_message(uid, "🎉 تمت ترقيتك إلى رتبة مطور بنجاح!")
                except Exception:
                    pass
            elif data.startswith("reject_"):
                uid = int(data[7:])
                if not user_is_super:
                    return
                if uid in db_data.get("pending_admins", []):
                    db_data["pending_admins"].remove(uid)
                save_db(db_data)
                bot.edit_message_text("❌ تم رفض الطلب.", chat_id=chat_id, message_id=msg_id)
            elif data == "back":
                bot.edit_message_text(
                    f"🌟 مرحباً {call.from_user.first_name}، اختر من القائمة:",
                    chat_id=chat_id,
                    message_id=msg_id,
                    reply_markup=main_menu(is_super=user_is_super)
                )
        except Exception as exc:
            logger.exception("Error processing callback data %s: %s", data, exc)

    def do_broadcast(msg, btype):
        db_data = load_db()
        ok_count = 0
        for uid in db_data.get("users", []):
            try:
                if btype == "text" and msg.text:
                    bot.send_message(uid, f"📢 إشعار:\n\n{msg.text}")
                elif btype == "photo" and msg.photo:
                    bot.send_photo(uid, msg.photo[-1].file_id, caption=msg.caption or "")
                elif btype == "video" and msg.video:
                    bot.send_video(uid, msg.video.file_id, caption=msg.caption or "")
                ok_count += 1
            except Exception:
                pass
        bot.reply_to(msg, f"✅ تم إرسال الإذاعة إلى {ok_count} مستخدم.")

    @bot.message_handler(content_types=["photo"])
    def handle_photo(message):
        remember_user(message.from_user)
        try:
            bot.send_chat_action(message.chat.id, "typing")
            photo = message.photo[-1]
            raw, _ = download_telegram_file(photo.file_id)
            prompt = message.caption or "صف الصورة بدقة وحلل محتواها بالعربية."
            reply = ask_vision(prompt, raw, "image/jpeg")
            send_long(message.chat.id, reply or "⚠️ لم أتمكن من استخراج تفاصيل الصورة.")
        except Exception as exc:
            bot.reply_to(message, f"⚠️ خطأ أثناء تحليل الصورة: {exc}")

    @bot.message_handler(content_types=["document"])
    def handle_document(message):
        remember_user(message.from_user)
        try:
            bot.send_chat_action(message.chat.id, "typing")
            filename = safe_filename(message.document.file_name or "file")
            raw, _ = download_telegram_file(message.document.file_id)
            bot.send_message(message.chat.id, f"📎 جاري تحليل `{filename}`...", parse_mode="Markdown")
            text, kind = extract_text_from_file(filename, raw)
            if not text:
                bot.send_message(message.chat.id, "⚠️ نوع الملف غير مدعوم للتحليل النصي.")
                return
            reply = ask_ai(f"حلل المحتوى التالي:\n\n{text[:MAX_TEXT_CHARS]}")
            send_long(message.chat.id, reply or "⚠️ تعذر التحليل.")
        except Exception as exc:
            bot.reply_to(message, f"⚠️ حدث خطأ: {exc}")

    @bot.message_handler(content_types=["text"])
    def handle_text(message):
        if message.text.startswith("/"):
            return
        remember_user(message.from_user)
        try:
            bot.send_chat_action(message.chat.id, "typing")
        except Exception:
            pass
        reply = ask_ai(message.text, conversation_id=f"tg-{message.from_user.id}")
        send_long(message.chat.id, reply or "⚠️ تعذر الحصول على رد من الذكاء الاصطناعي حالياً.")

# ============================================================
# Entry Point
# ============================================================
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port)
